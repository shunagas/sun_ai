#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
すん旅 経理帳簿ひな形を作成する。
複式簿記の仕訳帳から、貸借対照表・損益計算書・資金繰り表を数式で自動連動させる。

【重要】経理帳簿.xlsxは実データが入るためgit管理対象外（.gitignore参照）。
このスクリプトは出力先に既にファイルがあると実行を中止する（実データを誤って
上書き・消去しないため）。まっさらなテンプレートを作り直したい場合のみ、
既存の経理帳簿.xlsxを別名でバックアップしてから --force 付きで実行すること。
"""
import argparse
import os
import sys

import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(SCRIPT_DIR, "..", "経理帳簿.xlsx")

parser = argparse.ArgumentParser()
parser.add_argument("--force", action="store_true",
                     help="既存の経理帳簿.xlsxがあっても上書きする（実データが消えるので要注意）")
args = parser.parse_args()

if os.path.exists(OUT_PATH) and not args.force:
    print(f"エラー: {OUT_PATH} は既に存在します。実データが入っている可能性があるため上書きしません。")
    print("作り直す場合は、既存ファイルを別名でバックアップしてから --force を付けて実行してください。")
    sys.exit(1)

# ---------- 共通スタイル ----------
NAVY = "1F3864"
LIGHTBLUE = "D9E2F3"
YELLOW = "FFF2CC"
GRAY = "F2F2F2"
WHITE = "FFFFFF"

header_font = Font(name="Meiryo UI", size=11, bold=True, color=WHITE)
header_fill = PatternFill("solid", fgColor=NAVY)
subheader_font = Font(name="Meiryo UI", size=11, bold=True, color="1F3864")
subheader_fill = PatternFill("solid", fgColor=LIGHTBLUE)
normal_font = Font(name="Meiryo UI", size=11)
bold_font = Font(name="Meiryo UI", size=11, bold=True)
input_fill = PatternFill("solid", fgColor=YELLOW)
calc_fill = PatternFill("solid", fgColor=GRAY)
title_font = Font(name="Meiryo UI", size=16, bold=True, color=NAVY)
note_font = Font(name="Meiryo UI", size=10, italic=True, color="808080")

thin = Side(style="thin", color="BFBFBF")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

CUR_FMT = '#,##0"円"'
DATE_FMT = 'yyyy/m/d'

def style_header_row(ws, row, last_col):
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_all

def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = openpyxl.Workbook()
wb.remove(wb.active)

# =========================================================
# 1. 使い方
# =========================================================
ws = wb.create_sheet("使い方")
ws.sheet_properties.tabColor = NAVY
set_col_widths(ws, [3, 90])
ws["B2"] = "すん旅 経理帳簿の使い方"
ws["B2"].font = title_font

guide = [
    "",
    "【全体の仕組み】",
    "この帳簿は「仕訳帳」に取引を1件ずつ入力するだけで、貸借対照表（BS）・損益計算書（PL）・資金繰り表が自動で計算されます。",
    "自分で計算したり転記したりする必要はありません。黄色いセルだけ入力してください。",
    "",
    "【複式簿記のいちばん簡単な考え方】",
    "1つの取引を「借方（左）」と「貸方（右）」の2つの科目に分けて記録します。金額は必ず同じ額になります。",
    "例：現金で交通費1,200円を払った → 借方：旅費交通費 1,200円／貸方：現金 1,200円",
    "例：note記事の売上3,000円が銀行口座に入金された → 借方：普通預金（事業用） 3,000円／貸方：売上高（note） 3,000円",
    "迷ったら「お金が増えた側」「お金が減った側」を考え、勘定科目一覧シートの『通常残高』列を参考にしてください。",
    "",
    "【シートの説明】",
    "① 勘定科目一覧：使う科目の一覧です。増やしたい科目があれば、この表の下に行を追加して構いません。",
    "② 仕訳帳：ここに取引を入力します。日付・借方科目・借方金額・貸方科目・貸方金額・摘要を入れてください。",
    "　　借方科目／貸方科目はセルをクリックするとプルダウンで選べます。",
    "　　『チェック』列に「金額不一致」と出た場合、借方金額と貸方金額が違うので入力ミスがないか確認してください。",
    "③ 科目別残高集計：仕訳帳から自動集計される裏側の表です。基本的に見るだけで、入力は不要です。",
    "④ 貸借対照表：ある時点での財産の状況（資産・負債・純資産）です。すべての入力を反映した「今の残高」を表示します。",
    "　　いちばん下の『検算』が0円になっていれば、貸借（左右）が一致しており正しく記帳できています。0円でなければ入力ミスがあります。",
    "⑤ 損益計算書：1年間の儲け（収益－費用＝当期純利益）です。青色申告決算書を作るときのベースになります。",
    "⑥ 資金繰り表：現金・預金が月ごとにどれだけ増減したかの一覧です（正式な税務書類ではなく、資金管理用の参考表です）。",
    "",
    "【青色申告について】",
    "65万円の青色申告特別控除を受けるには、複式簿記での記帳に加えて、貸借対照表・損益計算書の添付、期限内申告、",
    "e-Taxでの申告または優良な電子帳簿保存のいずれかが必要です（制度は変わることがあるため、確定申告前に最新情報を確認してください）。",
    "詳しいスケジュールは sun_accounting/materials/tax-schedule.md にまとめています。",
]
r = 3
for line in guide:
    ws.cell(row=r, column=2, value=line)
    if line.startswith("【"):
        ws.cell(row=r, column=2).font = subheader_font
    else:
        ws.cell(row=r, column=2).font = normal_font
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

# =========================================================
# 2. 勘定科目一覧
# =========================================================
ws = wb.create_sheet("勘定科目一覧")
ws.sheet_properties.tabColor = "2E75B6"
headers = ["科目コード", "勘定科目", "区分", "BS/PL区分", "通常残高", "備考"]
for i, h in enumerate(headers, start=1):
    ws.cell(row=1, column=i, value=h)
style_header_row(ws, 1, len(headers))
set_col_widths(ws, [10, 22, 10, 10, 10, 40])
ws.freeze_panes = "A2"

accounts = [
    # code, name, kubun, bspl, normal, note
    (101, "現金", "資産", "BS", "借方", "手元現金"),
    (102, "普通預金（事業用）", "資産", "BS", "借方", "事業用の銀行口座"),
    (103, "売掛金", "資産", "BS", "借方", "まだ入金されていない売上"),
    (104, "前払費用", "資産", "BS", "借方", "前払いした費用（例：年払いサブスク）"),
    (201, "未払金", "負債", "BS", "貸方", "まだ支払っていない経費"),
    (202, "未払費用", "負債", "BS", "貸方", "発生しているが未払いの費用"),
    (203, "借入金", "負債", "BS", "貸方", "借入がある場合に使用"),
    (301, "元入金", "純資産", "BS", "貸方", "個人事業の資本金にあたる科目"),
    (302, "事業主貸", "純資産", "BS", "借方", "事業のお金を生活費など事業外に使った額（資本のマイナス）"),
    (303, "事業主借", "純資産", "BS", "貸方", "生活費口座から事業に補填した額（資本のプラス）"),
    (401, "売上高（note）", "収益", "PL", "貸方", "note記事・マガジンの売上"),
    (402, "売上高（YouTube）", "収益", "PL", "貸方", "YouTube広告収益・メンバーシップ等"),
    (403, "雑収入", "収益", "PL", "貸方", "チップ・アフィリエイト等その他収入"),
    (501, "旅費交通費", "費用", "PL", "借方", "取材・撮影のための交通費"),
    (502, "宿泊費", "費用", "PL", "借方", "取材・撮影のための宿泊費"),
    (503, "通信費", "費用", "PL", "借方", "インターネット・電話代等"),
    (504, "消耗品費", "費用", "PL", "借方", "撮影機材の消耗品等（10万円未満）"),
    (505, "支払手数料", "費用", "PL", "借方", "振込手数料・決済手数料等"),
    (506, "広告宣伝費", "費用", "PL", "借方", "広告出稿費用"),
    (507, "新聞図書費", "費用", "PL", "借方", "書籍・雑誌・有料情報等"),
    (508, "会議費", "費用", "PL", "借方", "打ち合わせ費用"),
    (509, "接待交際費", "費用", "PL", "借方", "取材先への手土産等"),
    (510, "減価償却費", "費用", "PL", "借方", "高額機材（10万円以上）の減価償却費"),
    (511, "雑費", "費用", "PL", "借方", "上記に当てはまらない費用"),
]
ACC_FIRST_ROW = 2
ACC_LAST_ROW = ACC_FIRST_ROW + len(accounts) - 1  # 25

for i, (code, name, kubun, bspl, normal, note) in enumerate(accounts):
    row = ACC_FIRST_ROW + i
    ws.cell(row=row, column=1, value=code)
    ws.cell(row=row, column=2, value=name)
    ws.cell(row=row, column=3, value=kubun)
    ws.cell(row=row, column=4, value=bspl)
    ws.cell(row=row, column=5, value=normal)
    ws.cell(row=row, column=6, value=note)
    for c in range(1, 7):
        cell = ws.cell(row=row, column=c)
        cell.font = normal_font
        cell.border = border_all
        if c in (1, 5):
            cell.alignment = Alignment(horizontal="center")

# 区分ごとの行範囲（後続シートの数式で使う）
ASSET_RANGE = (2, 5)      # 101-104
LIAB_RANGE = (6, 8)       # 201-203
EQUITY_RANGE = (9, 11)    # 301-303
REV_RANGE = (12, 14)      # 401-403
EXP_RANGE = (15, 25)      # 501-511

# =========================================================
# 3. 仕訳帳
# =========================================================
ws = wb.create_sheet("仕訳帳")
ws.sheet_properties.tabColor = "C00000"
JE_HEADERS = ["No", "日付", "借方科目", "借方金額", "貸方科目", "貸方金額", "摘要", "チェック", "現金預金増減(自動計算)"]
for i, h in enumerate(JE_HEADERS, start=1):
    ws.cell(row=1, column=i, value=h)
style_header_row(ws, 1, len(JE_HEADERS))
set_col_widths(ws, [6, 12, 20, 13, 20, 13, 36, 12, 22])
ws.freeze_panes = "A2"

JE_LAST_ROW = 501  # データ入力可能な最終行（500件分）

sample_rows = [
    (1, "2027/01/01", "普通預金（事業用）", 500000, "元入金", 500000, "【記入例】事業開始にあたり自己資金を事業用口座に入金"),
    (2, "2027/01/10", "旅費交通費", 1200, "現金", 1200, "【記入例】取材のための交通費（電車代）"),
    (3, "2027/01/15", "普通預金（事業用）", 3000, "売上高（note）", 3000, "【記入例】note有料記事の売上入金"),
    (4, "2027/01/20", "通信費", 5000, "普通預金（事業用）", 5000, "【記入例】インターネット回線使用料の引き落とし"),
    (5, "2027/01/25", "現金", 30000, "事業主借", 30000, "【記入例】生活費口座から事業用現金へ補填"),
]

import datetime as _dt

for r_idx, no, date_str, dr_acc, dr_amt, cr_acc, cr_amt, memo in [
    (2 + i, *row) for i, row in enumerate(sample_rows)
]:
    y, m, d = [int(x) for x in date_str.split("/")]
    ws.cell(row=r_idx, column=1, value=no)
    ws.cell(row=r_idx, column=2, value=_dt.date(y, m, d))
    ws.cell(row=r_idx, column=3, value=dr_acc)
    ws.cell(row=r_idx, column=4, value=dr_amt)
    ws.cell(row=r_idx, column=5, value=cr_acc)
    ws.cell(row=r_idx, column=6, value=cr_amt)
    ws.cell(row=r_idx, column=7, value=memo)

for row in range(2, JE_LAST_ROW + 1):
    ws.cell(row=row, column=2).number_format = DATE_FMT
    ws.cell(row=row, column=4).number_format = CUR_FMT
    ws.cell(row=row, column=6).number_format = CUR_FMT
    ws.cell(row=row, column=9).number_format = CUR_FMT
    # チェック列
    ws.cell(
        row=row, column=8,
        value=f'=IF(AND(D{row}="",F{row}=""),"",IF(D{row}=F{row},"OK","金額不一致"))'
    )
    # 現金預金増減（現金・普通預金（事業用）の増減を自動計算。振替は自動的に相殺される）
    ws.cell(
        row=row, column=9,
        value=(
            f'=IF(OR(C{row}="現金",C{row}="普通預金（事業用）"),D{row},0)'
            f'-IF(OR(E{row}="現金",E{row}="普通預金（事業用）"),F{row},0)'
        )
    )
    for c in range(1, 10):
        cell = ws.cell(row=row, column=c)
        cell.font = normal_font
        cell.border = border_all
        if c in (3, 4, 5, 6, 7):  # 入力してほしい列
            if row <= 6:
                pass
            else:
                cell.fill = input_fill
        if c in (8, 9):
            cell.fill = calc_fill

# 入力例行にも色をつける（薄め）
for row in range(2, 7):
    for c in (2, 3, 4, 5, 6, 7):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="E2EFDA")

# データ入力規則（プルダウン）
dv = DataValidation(
    type="list",
    formula1=f"=勘定科目一覧!$B${ACC_FIRST_ROW}:$B${ACC_LAST_ROW}",
    allow_blank=True,
    showDropDown=False,
)
ws.add_data_validation(dv)
dv.add(f"C2:C{JE_LAST_ROW}")
dv.add(f"E2:E{JE_LAST_ROW}")

note_row = JE_LAST_ROW + 2
ws.cell(row=note_row, column=2,
        value="※黄色いセルに入力してください。借方科目・貸方科目はプルダウンから選べます。500件まで入力できます（足りない場合は行を追加してください）。")
ws.cell(row=note_row, column=2).font = note_font

# =========================================================
# 4. 科目別残高集計
# =========================================================
ws = wb.create_sheet("科目別残高集計")
ws.sheet_properties.tabColor = GRAY
BAL_HEADERS = ["科目コード", "勘定科目", "区分", "通常残高", "借方合計", "貸方合計", "残高"]
for i, h in enumerate(BAL_HEADERS, start=1):
    ws.cell(row=1, column=i, value=h)
style_header_row(ws, 1, len(BAL_HEADERS))
set_col_widths(ws, [10, 22, 10, 10, 14, 14, 14])
ws.freeze_panes = "A2"

for row in range(ACC_FIRST_ROW, ACC_LAST_ROW + 1):
    ws.cell(row=row, column=1, value=f"=勘定科目一覧!A{row}")
    ws.cell(row=row, column=2, value=f"=勘定科目一覧!B{row}")
    ws.cell(row=row, column=3, value=f"=勘定科目一覧!C{row}")
    ws.cell(row=row, column=4, value=f"=勘定科目一覧!E{row}")
    ws.cell(row=row, column=5,
            value=f"=SUMIF(仕訳帳!$C$2:$C${JE_LAST_ROW},B{row},仕訳帳!$D$2:$D${JE_LAST_ROW})")
    ws.cell(row=row, column=6,
            value=f"=SUMIF(仕訳帳!$E$2:$E${JE_LAST_ROW},B{row},仕訳帳!$F$2:$F${JE_LAST_ROW})")
    ws.cell(row=row, column=7,
            value=f'=IF(D{row}="借方",E{row}-F{row},F{row}-E{row})')
    for c in range(1, 8):
        cell = ws.cell(row=row, column=c)
        cell.font = normal_font
        cell.border = border_all
        if c in (5, 6, 7):
            cell.number_format = CUR_FMT
        if c == 1 or c == 4:
            cell.alignment = Alignment(horizontal="center")

BR = {  # 区分名 -> 行範囲
    "asset": ASSET_RANGE, "liab": LIAB_RANGE, "equity": EQUITY_RANGE,
    "rev": REV_RANGE, "exp": EXP_RANGE,
}

def kv_block(ws, start_row, title, acc_range, col_label, col_value, minus=False):
    r = start_row
    ws.cell(row=r, column=col_label, value=title)
    ws.cell(row=r, column=col_label).font = subheader_font
    ws.cell(row=r, column=col_label).fill = subheader_fill
    ws.merge_cells(start_row=r, start_column=col_label, end_row=r, end_column=col_label + 1)
    r += 1
    first_data_row = r
    for acc_row in range(acc_range[0], acc_range[1] + 1):
        ws.cell(row=r, column=col_label, value=f"=科目別残高集計!B{acc_row}")
        ws.cell(row=r, column=col_value, value=f"=科目別残高集計!G{acc_row}")
        ws.cell(row=r, column=col_value).number_format = CUR_FMT
        for c in (col_label, col_value):
            ws.cell(row=r, column=c).font = normal_font
            ws.cell(row=r, column=c).border = border_all
        r += 1
    last_data_row = r - 1
    return first_data_row, last_data_row, r

# =========================================================
# 5. 損益計算書（先に作成し、当期純利益の行番号を貸借対照表から参照する）
# =========================================================
ws = wb.create_sheet("損益計算書")
ws.sheet_properties.tabColor = "2E75B6"
set_col_widths(ws, [4, 26, 16])
ws["B2"] = "損益計算書（記帳した全期間の収益・費用）"
ws["B2"].font = title_font
ws.merge_cells("B2:C2")

r = 4
label_col, value_col = 2, 3
first_r, last_r, r_after_rev = kv_block(ws, r, "【収益の部】", REV_RANGE, label_col, value_col)
rev_total_row = r_after_rev
ws.cell(row=rev_total_row, column=label_col, value="収益合計")
ws.cell(row=rev_total_row, column=label_col).font = bold_font
ws.cell(row=rev_total_row, column=value_col, value=f"=SUM(C{first_r}:C{last_r})")
ws.cell(row=rev_total_row, column=value_col).font = bold_font
ws.cell(row=rev_total_row, column=value_col).number_format = CUR_FMT
for c in (label_col, value_col):
    ws.cell(row=rev_total_row, column=c).border = border_all

r2 = rev_total_row + 2
first_e, last_e, r_after_exp = kv_block(ws, r2, "【費用の部】", EXP_RANGE, label_col, value_col)
exp_total_row = r_after_exp
ws.cell(row=exp_total_row, column=label_col, value="費用合計")
ws.cell(row=exp_total_row, column=label_col).font = bold_font
ws.cell(row=exp_total_row, column=value_col, value=f"=SUM(C{first_e}:C{last_e})")
ws.cell(row=exp_total_row, column=value_col).font = bold_font
ws.cell(row=exp_total_row, column=value_col).number_format = CUR_FMT
for c in (label_col, value_col):
    ws.cell(row=exp_total_row, column=c).border = border_all

profit_calc_row = exp_total_row + 2
ws.cell(row=profit_calc_row, column=label_col, value="当期純利益")
ws.cell(row=profit_calc_row, column=label_col).font = bold_font
ws.cell(row=profit_calc_row, column=value_col, value=f"=C{rev_total_row}-C{exp_total_row}")
ws.cell(row=profit_calc_row, column=value_col).font = bold_font
ws.cell(row=profit_calc_row, column=value_col).number_format = CUR_FMT
for c in (label_col, value_col):
    ws.cell(row=profit_calc_row, column=c).border = border_all
    ws.cell(row=profit_calc_row, column=c).fill = subheader_fill

PROFIT_ROW = profit_calc_row  # 貸借対照表シートから参照する行（C列）

# =========================================================
# 6. 貸借対照表
# =========================================================
ws = wb.create_sheet("貸借対照表")
ws.sheet_properties.tabColor = "2E75B6"
set_col_widths(ws, [4, 26, 16, 4, 26, 16])
ws["B2"] = "貸借対照表（すべての記帳を反映した現在時点の残高）"
ws["B2"].font = title_font
ws.merge_cells("B2:F2")

# 資産の部（左側）
r = 4
label_col, value_col = 2, 3
first, last, r_after_assets = kv_block(ws, r, "【資産の部】", ASSET_RANGE, label_col, value_col)
asset_total_row = r_after_assets
ws.cell(row=asset_total_row, column=label_col, value="資産合計")
ws.cell(row=asset_total_row, column=label_col).font = bold_font
ws.cell(row=asset_total_row, column=value_col,
        value=f"=SUM(C{first}:C{last})")
ws.cell(row=asset_total_row, column=value_col).font = bold_font
ws.cell(row=asset_total_row, column=value_col).number_format = CUR_FMT
for c in (label_col, value_col):
    ws.cell(row=asset_total_row, column=c).border = border_all

# 負債の部（右側）
r2 = 4
label_col2, value_col2 = 5, 6
first_l, last_l, r_after_liab = kv_block(ws, r2, "【負債の部】", LIAB_RANGE, label_col2, value_col2)
liab_total_row = r_after_liab
ws.cell(row=liab_total_row, column=label_col2, value="負債合計")
ws.cell(row=liab_total_row, column=label_col2).font = bold_font
ws.cell(row=liab_total_row, column=value_col2, value=f"=SUM(F{first_l}:F{last_l})")
ws.cell(row=liab_total_row, column=value_col2).font = bold_font
ws.cell(row=liab_total_row, column=value_col2).number_format = CUR_FMT
for c in (label_col2, value_col2):
    ws.cell(row=liab_total_row, column=c).border = border_all

# 純資産の部（右側、負債の下）
r3 = liab_total_row + 2
ws.cell(row=r3, column=label_col2, value="【純資産の部】")
ws.cell(row=r3, column=label_col2).font = subheader_font
ws.cell(row=r3, column=label_col2).fill = subheader_fill
ws.merge_cells(start_row=r3, start_column=label_col2, end_row=r3, end_column=label_col2 + 1)
r3 += 1
motoirekin_row = r3
ws.cell(row=r3, column=label_col2, value="元入金")
ws.cell(row=r3, column=value_col2, value="=科目別残高集計!G9")
r3 += 1
jigyounushikashi_row = r3
ws.cell(row=r3, column=label_col2, value="事業主貸（差引）")
ws.cell(row=r3, column=value_col2, value="=-科目別残高集計!G10")
r3 += 1
jigyounushikari_row = r3
ws.cell(row=r3, column=label_col2, value="事業主借")
ws.cell(row=r3, column=value_col2, value="=科目別残高集計!G11")
r3 += 1
profit_row = r3
ws.cell(row=r3, column=label_col2, value="当期純利益")
ws.cell(row=r3, column=value_col2, value=f"=損益計算書!C{PROFIT_ROW}")
r3 += 1
equity_total_row = r3
ws.cell(row=r3, column=label_col2, value="純資産合計")
ws.cell(row=r3, column=label_col2).font = bold_font
ws.cell(row=r3, column=value_col2,
        value=f"=F{motoirekin_row}+F{jigyounushikashi_row}+F{jigyounushikari_row}+F{profit_row}")
ws.cell(row=r3, column=value_col2).font = bold_font

for rr in range(r3 - 4, equity_total_row + 1):
    for c in (label_col2, value_col2):
        cell = ws.cell(row=rr, column=c)
        if not cell.font.bold:
            cell.font = normal_font
        cell.border = border_all
        if c == value_col2:
            cell.number_format = CUR_FMT

# 負債・純資産合計
grand_row = equity_total_row + 2
ws.cell(row=grand_row, column=label_col2, value="負債・純資産合計")
ws.cell(row=grand_row, column=label_col2).font = bold_font
ws.cell(row=grand_row, column=value_col2, value=f"=F{liab_total_row}+F{equity_total_row}")
ws.cell(row=grand_row, column=value_col2).font = bold_font
ws.cell(row=grand_row, column=value_col2).number_format = CUR_FMT
for c in (label_col2, value_col2):
    ws.cell(row=grand_row, column=c).border = border_all

# 検算
check_row = max(asset_total_row, grand_row) + 3
ws.cell(row=check_row, column=2, value="検算（資産合計 − 負債・純資産合計）")
ws.cell(row=check_row, column=2).font = bold_font
ws.cell(row=check_row, column=3, value=f"=C{asset_total_row}-F{grand_row}")
ws.cell(row=check_row, column=3).number_format = CUR_FMT
ws.cell(row=check_row, column=3).font = bold_font
ws.cell(row=check_row + 1, column=2,
        value="※ 0円になっていれば、貸借（左右）が一致しており正しく記帳できています。")
ws.cell(row=check_row + 1, column=2).font = note_font
ws.merge_cells(start_row=check_row + 1, start_column=2, end_row=check_row + 1, end_column=6)

# =========================================================
# 7. 資金繰り表（簡易キャッシュフロー、2027年分）
# =========================================================
ws = wb.create_sheet("資金繰り表")
ws.sheet_properties.tabColor = "70AD47"
set_col_widths(ws, [20] + [11] * 12 + [13])
ws["B2"] = "資金繰り表（現金・預金の月次増減：2027年）"
ws["B2"].font = title_font
ws.merge_cells("B2:N2")
ws["B3"] = "※ 正式な税務書類ではなく、資金繰りを把握するための参考表です（青色申告に必須なのは貸借対照表・損益計算書です）。"
ws["B3"].font = note_font
ws.merge_cells("B3:N3")

ws["B5"] = "期首残高（2027/1/1時点の現金・預金合計）"
ws["B5"].font = bold_font
ws.merge_cells("B5:E5")
ws["F5"] = 0
ws["F5"].number_format = CUR_FMT
ws["F5"].fill = input_fill
ws["F5"].border = border_all
OPENING_CELL = "F5"

HDR_ROW = 7
ws.cell(row=HDR_ROW, column=1, value="項目")
month_labels = [f"{m}月" for m in range(1, 13)]
for i, label in enumerate(month_labels):
    ws.cell(row=HDR_ROW, column=2 + i, value=label)
ws.cell(row=HDR_ROW, column=14, value="合計")
style_header_row(ws, HDR_ROW, 14)

FLOW_ROW = HDR_ROW + 1
CUM_ROW = HDR_ROW + 2
ws.cell(row=FLOW_ROW, column=1, value="当月収支（現金・預金増減）")
ws.cell(row=CUM_ROW, column=1, value="累計残高")
for r_ in (FLOW_ROW, CUM_ROW):
    ws.cell(row=r_, column=1).font = bold_font
    ws.cell(row=r_, column=1).border = border_all

for i in range(12):
    col = 2 + i
    col_letter = get_column_letter(col)
    # 当月収支：仕訳帳の現金預金増減(I列)を、該当月の日付だけ合計する
    ws.cell(
        row=FLOW_ROW, column=col,
        value=(
            f"=SUMIFS(仕訳帳!$I$2:$I${JE_LAST_ROW},"
            f"仕訳帳!$B$2:$B${JE_LAST_ROW},\">=\"&DATE(2027,{i+1},1),"
            f"仕訳帳!$B$2:$B${JE_LAST_ROW},\"<=\"&EOMONTH(DATE(2027,{i+1},1),0))"
        )
    )
    # 累計残高：1月は期首残高＋当月収支、2月以降は前月の累計残高＋当月収支
    if i == 0:
        ws.cell(row=CUM_ROW, column=col, value=f"={OPENING_CELL}+{col_letter}{FLOW_ROW}")
    else:
        prev_col_letter = get_column_letter(col - 1)
        ws.cell(row=CUM_ROW, column=col, value=f"={prev_col_letter}{CUM_ROW}+{col_letter}{FLOW_ROW}")
    for r_ in (FLOW_ROW, CUM_ROW):
        cell = ws.cell(row=r_, column=col)
        cell.number_format = CUR_FMT
        cell.font = normal_font
        cell.border = border_all
        cell.fill = calc_fill

# 合計列
ws.cell(row=FLOW_ROW, column=14, value=f"=SUM(B{FLOW_ROW}:M{FLOW_ROW})")
ws.cell(row=CUM_ROW, column=14, value=f"=M{CUM_ROW}")
for c in (14,):
    for r_ in (FLOW_ROW, CUM_ROW):
        cell = ws.cell(row=r_, column=c)
        cell.number_format = CUR_FMT
        cell.font = bold_font
        cell.border = border_all
        cell.fill = subheader_fill

note_row2 = CUM_ROW + 2
ws.cell(row=note_row2, column=2,
        value="※ 2027年分のひな形です。翌年以降は、このシートをコピーして年を書き換えて使ってください。")
ws.cell(row=note_row2, column=2).font = note_font
ws.merge_cells(start_row=note_row2, start_column=2, end_row=note_row2, end_column=8)

# =========================================================
# 保存
# =========================================================
wb.active = wb.sheetnames.index("使い方")
wb.save(OUT_PATH)
print("saved:", OUT_PATH)
print("sheets:", wb.sheetnames)
